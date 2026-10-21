import os
import sys
import re
import json
import zipfile
import xml.etree.ElementTree as ET
import docx

sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = os.path.dirname(os.path.abspath(__file__))
APP_DIR = os.path.join(ROOT_DIR, 'web_app')
DATA_DIR = os.path.join(APP_DIR, 'data')
ASSETS_DIR = os.path.join(APP_DIR, 'assets')
SLIDES_IMG_DIR = os.path.join(ASSETS_DIR, 'slides')

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(SLIDES_IMG_DIR, exist_ok=True)

def parse_kanji():
    kanji_path = os.path.join(ROOT_DIR, 'kanji', 'kanji.docx')
    doc = docx.Document(kanji_path)
    lessons = ['第 6 課', '第 7 課', '第 8 課', '第 9 課', '第 10 課']
    kanji_list = []

    for t_idx, t in enumerate(doc.tables):
        lesson = lessons[t_idx] if t_idx < len(lessons) else f'第 {t_idx+6} 課'
        r_idx = 0
        while r_idx < len(t.rows):
            r1 = [c.text.strip() for c in t.rows[r_idx].cells]
            r2 = [c.text.strip() for c in t.rows[r_idx+1].cells] if r_idx + 1 < len(t.rows) else ['', '', '']
            
            c0 = r1[0].replace('\n', ' ')
            m = re.search(r'([^\(\s]+)\s*\((.*?)\)', c0)
            kanji_char = m.group(1).strip() if m else c0
            han_viet = m.group(2).strip() if m else ''
            
            onyomi = r1[1].replace('音：', '').replace('音:', '').strip() if len(r1) > 1 else ''
            kunyomi = r1[2].replace('訓：', '').replace('訓:', '').strip() if len(r1) > 2 else ''
            
            examples_on = r2[1] if len(r2) > 1 else ''
            examples_kun = r2[2] if len(r2) > 2 else ''
            examples_raw = (examples_on + '\n' + examples_kun).strip()

            # Parse individual compound words
            compounds = []
            for line in re.split(r'[\n\|]+', examples_raw):
                line = line.strip()
                if not line: continue
                # Match pattern: 形成 (けいせい): hình thành
                cm = re.match(r'^(.*?)\s*\((.*?)\)\s*:\s*(.*)$', line)
                if cm:
                    compounds.append({
                        'word': cm.group(1).strip(),
                        'reading': cm.group(2).strip(),
                        'meaning': cm.group(3).strip()
                    })
                else:
                    compounds.append({
                        'word': line,
                        'reading': '',
                        'meaning': ''
                    })
            
            kanji_list.append({
                'id': f'kanji_{len(kanji_list)+1}',
                'lesson': lesson,
                'kanji': kanji_char,
                'han_viet': han_viet,
                'onyomi': onyomi,
                'kunyomi': kunyomi,
                'examples_raw': examples_raw,
                'compounds': compounds
            })
            r_idx += 2

    print(f"Parsed {len(kanji_list)} Kanji items.")
    return kanji_list

def parse_vocab():
    vocab_path = os.path.join(ROOT_DIR, 'tu_vung', 'toan_bo_tu_vung.docx')
    doc = docx.Document(vocab_path)
    current_lesson = '第 6 課'
    current_part = 'Part 1'
    vocab_list = []

    for child in doc.element.body:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(child, doc)
            txt = p.text.strip()
            if not txt: continue
            if '第' in txt and '課' in txt and 'TỔNG HỢP' not in txt:
                current_lesson = txt
            elif 'Part' in txt or '知って楽しむ' in txt:
                current_part = txt.rstrip(':')
        elif tag == 'tbl':
            t = docx.table.Table(child, doc)
            for r in t.rows:
                c = [cell.text.strip() for cell in r.cells]
                if len(c) >= 3:
                    kanji_word = c[0]
                    kana_reading = c[1]
                    meaning = c[2]
                    if not kanji_word and not kana_reading and not meaning:
                        continue
                    display_word = kanji_word if kanji_word else kana_reading
                    vocab_list.append({
                        'id': f'vocab_{len(vocab_list)+1}',
                        'lesson': current_lesson,
                        'part': current_part,
                        'kanji': kanji_word,
                        'kana': kana_reading,
                        'display': display_word,
                        'meaning': meaning
                    })
    print(f"Parsed {len(vocab_list)} Vocabulary items.")
    return vocab_list

def parse_grammar():
    grammar_path = os.path.join(ROOT_DIR, 'ngu_phap', 'toan_bo_ngu_phap.docx')
    doc = docx.Document(grammar_path)
    current_lesson = '第 6 課'
    grammar_list = []
    curr_grammar = None
    in_examples = False

    for child in doc.element.body:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(child, doc)
            txt = p.text.strip()
            if not txt: continue
            if '第' in txt and '課' in txt and 'TỔNG HỢP' not in txt:
                current_lesson = txt
                continue
            
            # Check for grammar rule title
            is_rule_title = False
            m = re.match(r'^(\d+)\s*[\.\:](.*)', txt)
            if m:
                rest = m.group(2).strip()
                has_vietnamese = bool(re.search(r'[àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđ]', rest, re.I))
                if has_vietnamese or rest.endswith(':') or rest.endswith('：') or '~' in rest or '…' in rest:
                    is_rule_title = True
            
            if is_rule_title:
                if curr_grammar:
                    grammar_list.append(curr_grammar)
                in_examples = False
                curr_grammar = {
                    'id': f'grammar_{len(grammar_list)+1}',
                    'lesson': current_lesson,
                    'title': txt,
                    'formula': '',
                    'points': '',
                    'examples': []
                }
            elif curr_grammar:
                if txt.startswith('ポイント') or 'ポイント:' in txt or 'ポイント：' in txt:
                    in_examples = False
                    curr_grammar['points'] = (curr_grammar['points'] + '\n' + txt).strip()
                elif txt.startswith('例:') or txt.startswith('例：') or txt.startswith('例'):
                    in_examples = True
                    curr_grammar['examples'].append(txt)
                elif in_examples:
                    curr_grammar['examples'].append(txt)
                else:
                    curr_grammar['points'] = (curr_grammar['points'] + '\n' + txt).strip()

        elif tag == 'tbl':
            t = docx.table.Table(child, doc)
            tbl_text = '\n'.join(' | '.join(c.text.strip() for c in r.cells) for r in t.rows if any(c.text.strip() for c in r.cells))
            if curr_grammar:
                if not curr_grammar['formula']:
                    curr_grammar['formula'] = tbl_text
                else:
                    curr_grammar['formula'] += '\n' + tbl_text

    if curr_grammar:
        grammar_list.append(curr_grammar)

    print(f"Parsed {len(grammar_list)} Grammar points.")
    return grammar_list

def parse_slides():
    slide_dir = os.path.join(ROOT_DIR, 'slide')
    slide_files = [
        ('第 6 課', '第6課ー地図を広げる.pptx', '地図を広げる'),
        ('第 7 課', '第7課ー世代を超えた交流 .pptx', '世代を超えた交流'),
        ('第 8 課', '第8課ー世代を超えた交流.pptx', '気持ちを伝える'),
        ('第 9 課', '第9課 言葉を楽しむ.pptx', '言葉を楽しむ')
    ]
    all_slide_data = []

    for lesson_name, file_name, topic in slide_files:
        pptx_path = os.path.join(slide_dir, file_name)
        if not os.path.exists(pptx_path):
            print(f"Warning: {file_name} does not exist!")
            continue

        lesson_slug = lesson_name.replace(' ', '_')
        lesson_img_dir = os.path.join(SLIDES_IMG_DIR, lesson_slug)
        os.makedirs(lesson_img_dir, exist_ok=True)

        with zipfile.ZipFile(pptx_path, 'r') as z:
            # 1. Read presentation rels to map rId -> target slide
            rels_xml = z.read('ppt/_rels/presentation.xml.rels')
            r_tree = ET.fromstring(rels_xml)
            rel_map = {node.attrib['Id']: node.attrib['Target'] for node in r_tree}

            # 2. Read presentation.xml to get canonical slide order
            pres_xml = z.read('ppt/presentation.xml')
            p_tree = ET.fromstring(pres_xml)
            ordered_slide_targets = []
            for elem in p_tree.iter():
                if elem.tag.endswith('sldId'):
                    r_id = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    if r_id in rel_map:
                        target = rel_map[r_id]
                        if not target.startswith('ppt/'):
                            target = 'ppt/' + target
                        ordered_slide_targets.append(target)

            lesson_slides = []
            for s_idx, slide_file in enumerate(ordered_slide_targets):
                slide_num = s_idx + 1
                if slide_file not in z.namelist():
                    continue

                # Parse slide xml
                s_xml = z.read(slide_file)
                s_tree = ET.fromstring(s_xml)
                
                # Extract text lines by paragraphs
                paragraphs = []
                for p_elem in s_tree.iter():
                    if p_elem.tag.endswith('}p'):
                        txt = ''.join(t.text for t in p_elem.iter() if t.tag.endswith('}t') and t.text)
                        txt = txt.strip()
                        if txt:
                            paragraphs.append(txt)

                # Extract relationships for images
                base_name = os.path.basename(slide_file)
                slide_rel_path = f'ppt/slides/_rels/{base_name}.rels'
                slide_images = []
                
                if slide_rel_path in z.namelist():
                    rel_tree = ET.fromstring(z.read(slide_rel_path))
                    img_idx = 1
                    for r in rel_tree:
                        if 'image' in r.attrib.get('Type', ''):
                            target_img = r.attrib['Target'].replace('../', 'ppt/')
                            if target_img in z.namelist():
                                ext = os.path.splitext(target_img)[1] or '.png'
                                out_img_name = f'slide_{slide_num}_img_{img_idx}{ext}'
                                out_img_path = os.path.join(lesson_img_dir, out_img_name)
                                
                                # Write image file if not already extracted
                                if not os.path.exists(out_img_path):
                                    with open(out_img_path, 'wb') as img_out:
                                        img_out.write(z.read(target_img))
                                
                                rel_web_path = f'assets/slides/{lesson_slug}/{out_img_name}'
                                slide_images.append(rel_web_path)
                                img_idx += 1

                slide_title = paragraphs[0] if paragraphs else f"Slide {slide_num}"
                lesson_slides.append({
                    'slide_number': slide_num,
                    'title': slide_title,
                    'text_content': paragraphs,
                    'images': slide_images
                })

            all_slide_data.append({
                'lesson': lesson_name,
                'topic': topic,
                'file_name': file_name,
                'total_slides': len(lesson_slides),
                'slides': lesson_slides
            })
            print(f"Extracted {len(lesson_slides)} slides for {lesson_name} ({topic})")

    return all_slide_data

def main():
    print("=== Starting Complete Dataset Extraction ===")
    kanji = parse_kanji()
    vocab = parse_vocab()
    grammar = parse_grammar()
    slides = parse_slides()

    # Save individual JSON files
    with open(os.path.join(DATA_DIR, 'kanji.json'), 'w', encoding='utf-8') as f:
        json.dump(kanji, f, ensure_ascii=False, indent=2)

    with open(os.path.join(DATA_DIR, 'vocabulary.json'), 'w', encoding='utf-8') as f:
        json.dump(vocab, f, ensure_ascii=False, indent=2)

    with open(os.path.join(DATA_DIR, 'grammar.json'), 'w', encoding='utf-8') as f:
        json.dump(grammar, f, ensure_ascii=False, indent=2)

    with open(os.path.join(DATA_DIR, 'slides.json'), 'w', encoding='utf-8') as f:
        json.dump(slides, f, ensure_ascii=False, indent=2)

    # Save combined JS file so that index.html can load data via script tag
    # without any CORS restrictions if opened directly as file://
    combined_data = {
        'kanji': kanji,
        'vocabulary': vocab,
        'grammar': grammar,
        'slides': slides
    }
    
    js_content = f"// Auto-generated data for Nihongo Speed Master\nwindow.APP_DATA = {json.dumps(combined_data, ensure_ascii=False, indent=2)};\n"
    with open(os.path.join(DATA_DIR, 'app_data.js'), 'w', encoding='utf-8') as f:
        f.write(js_content)

    print("=== Extraction Finished Successfully! ===")
    print(f"Kanji count: {len(kanji)}")
    print(f"Vocabulary count: {len(vocab)}")
    print(f"Grammar count: {len(grammar)}")
    total_slides = sum(s['total_slides'] for s in slides)
    print(f"Total Slides extracted: {total_slides}")

if __name__ == '__main__':
    main()
