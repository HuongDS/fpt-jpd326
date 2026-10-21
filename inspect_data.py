import sys
import os
import docx
import json
import zipfile
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

def inspect_kanji():
    doc = docx.Document('kanji/kanji.docx')
    print("=== KANJI ===")
    print("Paragraphs:", [p.text.strip() for p in doc.paragraphs if p.text.strip()])
    print("Table count:", len(doc.tables))
    for t_idx, t in enumerate(doc.tables):
        print(f"\n--- Table {t_idx} (rows: {len(t.rows)}) ---")
        for r_idx in range(min(6, len(t.rows))):
            cells = [c.text.strip().replace('\n', ' | ') for c in t.rows[r_idx].cells]
            print(f"Row {r_idx}: {cells}")

def inspect_tu_vung():
    doc = docx.Document('tu_vung/toan_bo_tu_vung.docx')
    print("\n=== TU VUNG ===")
    print("Paragraphs count:", len([p.text.strip() for p in doc.paragraphs if p.text.strip()]))
    for p in doc.paragraphs[:20]:
        if p.text.strip():
            print("P:", p.text.strip())
    print("Table count:", len(doc.tables))
    for t_idx in range(min(5, len(doc.tables))):
        t = doc.tables[t_idx]
        print(f"\n--- Table {t_idx} (rows: {len(t.rows)}) ---")
        for r_idx in range(min(4, len(t.rows))):
            cells = [c.text.strip().replace('\n', ' | ') for c in t.rows[r_idx].cells]
            print(f"Row {r_idx}: {cells}")

def inspect_ngu_phap():
    doc = docx.Document('ngu_phap/toan_bo_ngu_phap.docx')
    print("\n=== NGU PHAP ===")
    print("Paragraphs count:", len(doc.paragraphs))
    for i, p in enumerate(doc.paragraphs[:35]):
        if p.text.strip():
            print(f"P{i}: {p.text.strip()}")
    print("Table count:", len(doc.tables))
    for t_idx in range(min(5, len(doc.tables))):
        t = doc.tables[t_idx]
        print(f"\n--- Table {t_idx} (rows: {len(t.rows)}) ---")
        for r_idx in range(len(t.rows)):
            cells = [c.text.strip().replace('\n', ' | ') for c in t.rows[r_idx].cells]
            print(f"Row {r_idx}: {cells}")

if __name__ == '__main__':
    inspect_kanji()
    inspect_tu_vung()
    inspect_ngu_phap()
